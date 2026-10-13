"""Отчёт: выполняет SQL-запросы, сохраняет report.xlsx (лист на запрос) и печатает
короткую сводку в консоль. Запуск: python -m pipeline.report
"""

from __future__ import annotations

from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import text

from pipeline.config import REPORT_PATH, SQL_DIR
from pipeline.db import engine

SHEETS = {
    "01_revenue_monthly_cumulative": "Выручка по месяцам",
    "02_rfm_segments": "RFM-сегменты",
    "03_warehouse_with_kit_expansion": "Склад (раскрытие)",
    "04_days_of_stock": "Остатки и дефицит",
    "05_repeat_purchases": "Повторные покупки",
    "06_top10_products_90d": "Топ-10 за 90 дней",
}

COLS = {
    "month": "Месяц", "revenue": "Выручка, ₽", "orders": "Заказов",
    "revenue_cumulative": "Нарастающий итог, ₽", "name": "Клиент",
    "phone_normalized": "Телефон", "last_order": "Последний заказ",
    "frequency": "Заказов", "monetary": "Сумма, ₽", "r": "R", "f": "F", "m": "M",
    "segment": "Сегмент", "sku": "Артикул", "title": "Наименование", "type": "Тип",
    "units_sold": "Продано, шт", "stock": "Остаток", "per_day": "Продаж/день",
    "days_left": "Хватит на, дней", "customers": "Клиентов",
    "repeat_customers": "С повтором", "repeat_share_pct": "Доля повторных, %",
    "avg_days_between_orders": "Средний интервал, дней", "share_pct": "Доля, %",
}


def _run_sql(sql: str):
    with engine.connect() as conn:
        res = conn.execute(text(sql))
        return list(res.keys()), res.fetchall()


def build() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    results: dict[str, tuple[list[str], list]] = {}

    for path in sorted(SQL_DIR.glob("*.sql")):
        stem = path.stem
        cols, rows = _run_sql(path.read_text(encoding="utf-8"))
        results[stem] = (cols, rows)
        ws = wb.create_sheet(title=SHEETS.get(stem, stem)[:31])
        ws.append([COLS.get(c, c) for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([_cell(v) for v in row])
        for i, c in enumerate(cols, start=1):
            width = max(len(COLS.get(c, c)), *(len(str(_cell(r[i - 1]))) for r in rows[:50])) if rows else len(c)
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width + 2, 40)

    wb.save(REPORT_PATH)
    _print_summary(results)


def _cell(v):
    if isinstance(v, Decimal):
        return float(v)
    return v


def _usd_rate() -> float | None:
    with engine.connect() as conn:
        row = conn.execute(
            text("SELECT value/nominal FROM fx_rates WHERE char_code='USD' ORDER BY rate_date DESC LIMIT 1")
        ).first()
    return float(row[0]) if row else None


def _print_summary(results: dict[str, tuple[list[str], list]]) -> None:
    rev_cols, rev_rows = results["01_revenue_monthly_cumulative"]
    total_rev = sum(r[rev_cols.index("revenue")] for r in rev_rows) if rev_rows else 0
    last_month = rev_rows[-1] if rev_rows else None

    rep_cols, rep_rows = results["05_repeat_purchases"]
    rep = rep_rows[0] if rep_rows else None

    top_cols, top_rows = results["06_top10_products_90d"]
    top = top_rows[0] if top_rows else None

    stock_cols, stock_rows = results["04_days_of_stock"]
    soon = next((r for r in stock_rows if r[stock_cols.index("days_left")] is not None), None)

    usd = _usd_rate()

    print("\nСводка по отчёту:")
    print(f"  Выручка всего (без отмен): {total_rev:,.0f} ₽".replace(",", " "))
    if usd:
        print(f"  Курс ЦБ USD: {usd:.2f} ₽ → выручка ≈ ${total_rev / usd:,.0f}".replace(",", " "))
    if last_month:
        print(f"  Последний месяц: {last_month[rev_cols.index('month')]} — "
              f"{last_month[rev_cols.index('revenue')]:,.0f} ₽".replace(",", " "))
    if rep:
        print(f"  Повторных клиентов: {rep[rep_cols.index('repeat_share_pct')]}% "
              f"(интервал ~{rep[rep_cols.index('avg_days_between_orders')]} дн.)")
    if top:
        print(f"  Топ-товар за 90 дней: {top[top_cols.index('title')]} — "
              f"{top[top_cols.index('revenue')]:,.0f} ₽, доля {top[top_cols.index('share_pct')]}%".replace(",", " "))
    if soon:
        print(f"  Кончится первым: {soon[stock_cols.index('title')]} — "
              f"хватит на {soon[stock_cols.index('days_left')]} дн.")
    print(f"  Файл: {REPORT_PATH.name}")


if __name__ == "__main__":
    build()
